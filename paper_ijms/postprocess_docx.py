"""Table borders/widths/keep-together, caption keep-with-next, post-table spacing."""
import sys
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
TEXT_W=16.0
def el(tag,**a):
    e=OxmlElement(tag)
    for k,v in a.items(): e.set(qn(f'w:{k}'),str(v))
    return e
def widths_for(t,font_pt):
    cw=0.021*font_pt; pad=0.35; mins=[]; prefs=[]
    for ci in range(len(t.columns)):
        cells=[r.cells[ci].text for r in t.rows]
        lw=max((len(w) for c in cells for w in c.split()),default=3); mean=sum(len(c) for c in cells)/len(cells)
        mins.append(lw*cw+pad); prefs.append(max(lw*cw+pad, mean*cw*0.75+pad))
    tot=sum(prefs)
    if tot<=TEXT_W: return [p+(TEXT_W-tot)*p/tot for p in prefs]
    flex=[p-m for p,m in zip(prefs,mins)]; need=tot-TEXT_W
    if sum(flex)<=0: return [m*TEXT_W/sum(mins) for m in mins]
    k=min(1,need/sum(flex)); return [p-f*k for p,f in zip(prefs,flex)]
HDR={"Q1 greedy CPU ± CI":"Q1 greedy CPU ± CI","Q1 greedy joint sat.":"Q1 greedy sat.","Q1 stochastic CPU ± CI":"Q1 stoch. CPU ± CI","Q1 stochastic joint sat.":"Q1 stoch. sat.","Q2 stochastic CPU ± CI":"Q2 stoch. CPU ± CI","Q2 joint sat.":"Q2 sat.","Q2 val.-satisfied":"Q2 val.-sat.","Val.-satisfied":"Val.-sat.","Stochastic joint sat.":"Stoch. joint sat.","Greedy joint sat.":"Greedy joint sat.","Stochastic CPU ± CI":"Stoch. CPU ± CI","Greedy CPU ± CI":"Greedy CPU ± CI","Stochastic TP":"Stoch. TP"}
d=Document(sys.argv[1])
for t in d.tables:
    n=len(t.columns); font=9.5 if n<=6 else (9 if n<=8 else 8.5)
    for cell in t.rows[0].cells:
        for par in cell.paragraphs:
            for run in par.runs:
                if run.text in HDR: run.text=HDR[run.text]
    tblPr=t._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')): tblPr.remove(old)
    b=el('w:tblBorders')
    for edge,val in [('top','single'),('bottom','single'),('insideH','single'),('left','nil'),('right','nil'),('insideV','nil')]:
        b.append(el(f'w:{edge}',val=val,sz=6,space=0,color='000000'))
    tblPr.append(b)
    w=widths_for(t,font)
    lay=tblPr.find(qn('w:tblLayout'))
    if lay is None: lay=el('w:tblLayout'); tblPr.append(lay)
    lay.set(qn('w:type'),'fixed')
    tw=tblPr.find(qn('w:tblW'))
    if tw is None: tw=el('w:tblW'); tblPr.append(tw)
    tw.set(qn('w:type'),'dxa'); tw.set(qn('w:w'),str(int(sum(w)*567)))
    for gc,wc in zip(t._tbl.tblGrid.findall(qn('w:gridCol')),w): gc.set(qn('w:w'),str(int(wc*567)))
    for ri,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr(); trPr.append(el('w:cantSplit'))
        if ri==0: trPr.append(el('w:tblHeader'))
        for cell,wc in zip(row.cells,w):
            cell.width=Cm(wc)
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(1); par.paragraph_format.space_before=Pt(1)
                par.paragraph_format.keep_with_next=(ri<len(t.rows)-1)
                for run in par.runs:
                    run.font.size=Pt(font); run.font.name='Times New Roman'
                    if ri==0: run.font.bold=True
    for cell in t.rows[0].cells:
        tcPr=cell._tc.get_or_add_tcPr(); tb=el('w:tcBorders'); tb.append(el('w:bottom',val='single',sz=12,color='000000')); tcPr.append(tb)
kids=list(d.element.body.iterchildren())
for i,k in enumerate(kids):
    if k.tag==qn('w:tbl'):
        if i>0 and kids[i-1].tag==qn('w:p'): Paragraph(kids[i-1],d).paragraph_format.keep_with_next=True
        if i+1<len(kids) and kids[i+1].tag==qn('w:p'): Paragraph(kids[i+1],d).paragraph_format.space_before=Pt(10)

# ---- submission furniture: centred page numbers and continuous line numbers ----
# Publishers ask for both so that reviewers can cite a location.  Line numbers
# restart nowhere (continuous) and count every line.
sect=d.sections[0]
sectPr=sect._sectPr
for old in sectPr.findall(qn('w:lnNumType')): sectPr.remove(old)
ln=el('w:lnNumType',countBy=1,start=1,restart='continuous')
# w:lnNumType must precede w:pgNumType/w:cols etc.; append is tolerated by Word,
# but insert after page margins to keep the schema order valid.
ref=sectPr.find(qn('w:pgMar'))
(ref.addnext(ln) if ref is not None else sectPr.append(ln))

footer=sect.footer
footer.is_linked_to_previous=False
p=footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
p.text=''
from docx.enum.text import WD_ALIGN_PARAGRAPH
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
def field(par,instr):
    r=par.add_run()._r
    f1=el('w:fldChar',fldCharType='begin'); r.append(f1)
    it=OxmlElement('w:instrText'); it.set(qn('xml:space'),'preserve'); it.text=instr; r.append(it)
    r.append(el('w:fldChar',fldCharType='separate'))
    t=OxmlElement('w:t'); t.text='1'; r.append(t)
    r.append(el('w:fldChar',fldCharType='end'))
field(p,' PAGE ')
for run in p.runs:
    run.font.size=Pt(10); run.font.name='Times New Roman'

d.save(sys.argv[1])
